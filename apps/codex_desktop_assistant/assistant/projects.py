"""Durable dependency scheduler above the chat/browser controllers.

Safe local recipes run in bounded workers. Account operations require a registered
adapter and an owner-approved immutable batch manifest. An uncertain external
attempt is reconciled before retry; it is never silently replayed after restart.
"""
import concurrent.futures
import csv
import hashlib
import io
import json
import sqlite3
import threading
import time
import uuid
from pathlib import Path

from .permissions import Denied
from .products import ProductFactory,RECIPES,DOCUMENT_RECIPES,CUSTOM_RECIPES,digest

SAFE={'generate','validate','package'}
KINDS=SAFE|{'publish','verify','reason'}
MAX_JOBS=20000


def encoded(value):return json.dumps(value,sort_keys=True,ensure_ascii=False,allow_nan=False,separators=(',',':'))
def ident():return uuid.uuid4().hex


class ProjectQueue:
    def __init__(self,directory,emit=lambda *args:None,workers=2):
        self.root=Path(directory)/'projects';self.root.mkdir(parents=True,exist_ok=True)
        self.lock=threading.RLock();self.wake=threading.Event();self.closed=threading.Event();self.emit=emit
        self.workers=max(1,min(int(workers),4));self.adapters={};self.futures={};self.dispatch=None
        self.db=sqlite3.connect(self.root/'queue.sqlite3',check_same_thread=False);self.db.row_factory=sqlite3.Row
        self.db.executescript('''PRAGMA journal_mode=WAL; PRAGMA foreign_keys=ON;
        CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY,conversation TEXT,objective TEXT,state TEXT,version INTEGER DEFAULT 1,created REAL,updated REAL);
        CREATE TABLE IF NOT EXISTS jobs(id TEXT PRIMARY KEY,project TEXT REFERENCES projects(id),key TEXT,kind TEXT,state TEXT,payload TEXT,result TEXT DEFAULT '{}',priority INTEGER DEFAULT 0,attempts INTEGER DEFAULT 0,max_attempts INTEGER DEFAULT 3,next_at REAL DEFAULT 0,error TEXT DEFAULT '',conversation TEXT,version INTEGER,UNIQUE(project,key));
        CREATE TABLE IF NOT EXISTS dependencies(job TEXT REFERENCES jobs(id),needs TEXT REFERENCES jobs(id),PRIMARY KEY(job,needs));
        CREATE TABLE IF NOT EXISTS events(seq INTEGER PRIMARY KEY AUTOINCREMENT,project TEXT,job TEXT,kind TEXT,detail TEXT,at REAL);
        CREATE TABLE IF NOT EXISTS scopes(id TEXT PRIMARY KEY,project TEXT,store TEXT,manifest TEXT,expires REAL,revoked INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS effects(key TEXT PRIMARY KEY,project TEXT,job TEXT,store TEXT,state TEXT,fingerprint TEXT,result TEXT,updated REAL);
        CREATE TABLE IF NOT EXISTS reports(project TEXT PRIMARY KEY,state TEXT);
        CREATE INDEX IF NOT EXISTS ready_jobs ON jobs(state,next_at,priority);
        ''')
        # Recovery is conservative even when an adapter supports idempotency.
        self.db.execute("UPDATE jobs SET state=CASE WHEN kind IN ('generate','validate','package') THEN 'pending' ELSE 'review' END,error='Interrupted process: inspect checkpoint before consequential continuation' WHERE state IN ('running','waiting_ai')")
        self.db.commit()
        self.pool=concurrent.futures.ThreadPoolExecutor(max_workers=self.workers,thread_name_prefix='project-worker')
        self.thread=None

    def rows(self,sql,args=()):
        with self.lock:return [dict(r) for r in self.db.execute(sql,args)]

    def event(self,project,job,kind,detail):
        self.db.execute('INSERT INTO events(project,job,kind,detail,at) VALUES(?,?,?,?,?)',(project,job,kind,encoded(detail),time.time()))

    def owned(self,project,conversation):
        rows=self.rows('SELECT * FROM projects WHERE id=? AND conversation=?',(project,conversation))
        if not rows:raise Denied('Project belongs to another chat or does not exist')
        return rows[0]

    def create(self,conversation,objective):
        if not isinstance(objective,str) or not objective.strip() or len(objective)>100000:raise ValueError('Provide a bounded project objective')
        project=ident();now=time.time()
        with self.lock:
            if self.db.execute('SELECT COUNT(*) FROM projects').fetchone()[0]>=1000:raise ValueError('Local project quota reached; archive/export old projects before adding more')
            self.db.execute('INSERT INTO projects VALUES(?,?,?,?,?,?,?)',(project,conversation,objective,'draft',1,now,now))
            self.event(project,None,'created',{'objective':objective});self.db.commit()
        return {'project':project,'state':'draft','workspace':str(self.root/project/'artifacts')}

    def add(self,project,conversation,jobs,importing=False):
        self.owned(project,conversation)
        if not isinstance(jobs,list) or not 1<=len(jobs)<=(MAX_JOBS if importing else 200):raise ValueError('Add 1–200 tasks per batch; additional batches extend the same durable queue')
        with self.lock:
            if self.db.execute('SELECT COUNT(*) FROM jobs WHERE project=?',(project,)).fetchone()[0]+len(jobs)>MAX_JOBS:raise ValueError('Project task quota exceeded')
            existing={r['key']:r for r in self.db.execute('SELECT * FROM jobs WHERE project=?',(project,))};incoming={}
            for job in jobs:
                key=job.get('key');kind=job.get('kind');payload=job.get('payload',{});needs=job.get('needs',[])
                if not isinstance(key,str) or not key or len(key)>160 or kind not in KINDS or len(encoded(payload))>20000 or not isinstance(needs,list) or len(needs)>200:raise ValueError('Invalid bounded task')
                if kind in SAFE|{'publish','verify'}:
                    product=payload.get('product',{})
                    key_id=product.get('id','')
                    if not key_id or len(key_id)>80 or any(c not in 'abcdefghijklmnopqrstuvwxyz0123456789-_' for c in key_id) or product.get('recipe') not in (set(RECIPES)|DOCUMENT_RECIPES|CUSTOM_RECIPES):raise ValueError('Invalid recipe product identity')
                    if len(str(product.get('title','')))>120 or type(product.get('price_cents',0)) is not int or not 0<=product.get('price_cents',0)<=10000000:raise ValueError('Invalid recipe metadata')
                signature=(kind,encoded(payload),sorted(set(needs)))
                if key in incoming and incoming[key][0]!=signature:raise ValueError('Conflicting duplicate task key')
                if key in existing:
                    old=existing[key];old_needs=[r[0] for r in self.db.execute('SELECT j.key FROM dependencies d JOIN jobs j ON j.id=d.needs WHERE d.job=?',(old['id'],))]
                    if (old['kind'],old['payload'],sorted(old_needs))!=signature:raise ValueError('Task key already has different content; completed work is immutable')
                incoming[key]=(signature,job)
            known=set(existing)|set(incoming)
            graph={key:[r[0] for r in self.db.execute('SELECT j.key FROM dependencies d JOIN jobs j ON j.id=d.needs WHERE d.job=?',(row['id'],))] for key,row in existing.items()}
            graph.update({key:entry[0][2] for key,entry in incoming.items()})
            # Iterative topological check handles thousands of jobs without recursion.
            for needs in graph.values():
                if any(n not in known for n in needs):raise ValueError('Unknown dependency')
            done=set();remaining=dict(graph)
            while remaining:
                ready={k for k,v in remaining.items() if set(v)<=done}
                if not ready:raise ValueError('Dependency cycle')
                done|=ready
                for key in ready:remaining.pop(key)
            version=self.db.execute('SELECT version FROM projects WHERE id=?',(project,)).fetchone()[0]
            try:
                inserted=False
                for key,(signature,job) in incoming.items():
                    if key in existing:continue
                    inserted=True
                    row_id=ident();self.db.execute('INSERT INTO jobs(id,project,key,kind,state,payload,priority,version) VALUES(?,?,?,?,?,?,?,?)',(row_id,project,key,job['kind'],'pending',signature[1],max(-100,min(100,int(job.get('priority',0)))),version));existing[key]={'id':row_id}
                for key,(signature,_) in incoming.items():
                    for dependency in signature[2]:self.db.execute('INSERT OR IGNORE INTO dependencies VALUES(?,?)',(existing[key]['id'],existing[dependency]['id']))
                if inserted:
                    self.db.execute("UPDATE projects SET state='draft' WHERE id=? AND state='completed'",(project,));self.db.execute('DELETE FROM reports WHERE project=?',(project,))
                self.event(project,None,'plan_extended',{'tasks':len(incoming)});self.db.commit()
            except BaseException:self.db.rollback();raise
        self.wake.set();return self.summary(project,conversation)

    def products(self,project,conversation,products,store=None):
        if not isinstance(products,list) or not 1<=len(products)<=40:raise ValueError('Plan 1–40 products per batch')
        jobs=[];builtins={}
        for row in self.rows("SELECT payload FROM jobs WHERE project=? AND kind='generate'",(project,)):
            previous=json.loads(row['payload']).get('product',{})
            if previous.get('recipe') in RECIPES:builtins[previous['recipe']]=previous['id']
        for product in products:
            product=dict(product);key=product.get('id')
            if not isinstance(key,str) or not key or len(key)>80 or any(c not in 'abcdefghijklmnopqrstuvwxyz0123456789-_' for c in key):raise ValueError('Use a short lowercase product ID')
            if product.get('recipe') not in (set(RECIPES)|DOCUMENT_RECIPES|CUSTOM_RECIPES):raise ValueError('Choose a supported recipe or plan a reasoning task for a custom design')
            if product['recipe'] in RECIPES:
                if product['recipe'] in builtins and builtins[product['recipe']]!=key:raise ValueError('Relabeling the same built-in template is duplicate filler. Use a distinct custom spreadsheet or document specification.')
                builtins[product['recipe']]=key
            if len(str(product.get('title','')))>120 or type(product.get('price_cents',0)) is not int or not 0<=product.get('price_cents',0)<=10000000:raise ValueError('Invalid product title/price')
            for stage in ('generate','validate','package'):
                preceding={'generate':[],'validate':[key+':generate'],'package':[key+':generate',key+':validate']}[stage]
                jobs.append({'key':key+':'+stage,'kind':stage,'payload':{'product':product},'needs':preceding})
            if store:
                jobs.extend([{'key':key+':publish','kind':'publish','payload':{'product':product,'store':store},'needs':[key+':package']},{'key':key+':verify','kind':'verify','payload':{'product':product,'store':store},'needs':[key+':publish']}])
        return self.add(project,conversation,jobs)

    def summary(self,project,conversation,offset=0,limit=50):
        p=self.owned(project,conversation);limit=max(1,min(int(limit),100));offset=max(0,int(offset))
        counts={r['state']:r['count'] for r in self.rows('SELECT state,COUNT(*) count FROM jobs WHERE project=? GROUP BY state',(project,))}
        jobs=self.rows('SELECT id,key,kind,state,priority,attempts,error FROM jobs WHERE project=? ORDER BY rowid LIMIT ? OFFSET ?',(project,limit,offset))
        return {'project':project,'objective':p['objective'],'state':p['state'],'version':p['version'],'counts':counts,'total':sum(counts.values()),'jobs':jobs,'offset':offset,'next_offset':offset+len(jobs) if offset+len(jobs)<sum(counts.values()) else None}

    def control(self,project,conversation,operation):
        self.owned(project,conversation)
        if operation not in ('start','pause','cancel','resume'):raise ValueError('Unsupported project control')
        with self.lock:
            state={'start':'running','resume':'running','pause':'paused','cancel':'cancelled'}[operation]
            current=self.db.execute('SELECT state FROM projects WHERE id=?',(project,)).fetchone()[0]
            if current=='cancelled':raise Denied('Cancelled projects cannot resume; create a new plan referencing retained results')
            self.db.execute('UPDATE projects SET state=?,updated=? WHERE id=?',(state,time.time(),project))
            if operation=='cancel':
                self.db.execute("UPDATE jobs SET state='cancelled' WHERE project=? AND state NOT IN ('completed','failed')",(project,));self.db.execute('UPDATE scopes SET revoked=1 WHERE project=?',(project,))
            self.event(project,None,operation,{});self.db.commit()
        self.start();self.wake.set();return self.summary(project,conversation)

    def revise(self,project,conversation,objective):
        self.owned(project,conversation)
        if not isinstance(objective,str) or not objective.strip() or len(objective)>100000:raise ValueError('Invalid objective')
        with self.lock:
            self.db.execute("UPDATE projects SET objective=?,version=version+1,state='paused',updated=? WHERE id=?",(objective,time.time(),project))
            self.db.execute("UPDATE jobs SET state='review',error='Objective changed; retain completed work and review remaining scope' WHERE project=? AND state!='completed'",(project,));self.db.execute('UPDATE scopes SET revoked=1 WHERE project=?',(project,));self.event(project,None,'objective_changed',{});self.db.commit()
        return self.summary(project,conversation)

    def retry(self,project,conversation,job_id):
        p=self.owned(project,conversation)
        with self.lock:
            row=self.db.execute('SELECT * FROM jobs WHERE project=? AND id=?',(project,job_id)).fetchone()
            if not row or row['state'] not in ('failed','blocked','review'):raise ValueError('Choose a failed, blocked or review task')
            # Explicit owner/model request may requeue read-only reconciliation;
            # write handler still must reconcile effects and revalidate scope.
            self.db.execute("UPDATE jobs SET state='pending',attempts=0,next_at=0,error='',version=? WHERE id=?",(p['version'],job_id));self.event(project,job_id,'retry_requested',{});self.db.commit()
            self.db.execute("UPDATE projects SET state='running' WHERE id=? AND state='blocked'",(project,));self.db.commit()
        self.wake.set();return self.summary(project,conversation)

    def priority(self,project,conversation,job_id,value):
        self.owned(project,conversation)
        with self.lock:self.db.execute('UPDATE jobs SET priority=? WHERE id=? AND project=?',(max(-100,min(100,int(value))),job_id,project));self.db.commit()

    def result(self,project,conversation,key):
        self.owned(project,conversation);rows=self.rows('SELECT state,result,error FROM jobs WHERE project=? AND key=?',(project,key))
        if not rows:raise ValueError('Task not found')
        row=rows[0];row['result']=json.loads(row['result']);return row

    def import_attachment(self,project,conversation,record,attachments):
        self.owned(project,conversation);path=attachments.path(record)
        if digest(path)!=record['sha256']:raise Denied('Task attachment changed')
        if path.stat().st_size>10*1024*1024:raise ValueError('Task imports are limited to 10 MB')
        suffix=path.suffix.lower()
        if suffix=='.json':
            data=json.loads(path.read_text(encoding='utf-8-sig'));rows=data.get('tasks',data.get('products',[])) if isinstance(data,dict) else data
        elif suffix in ('.csv','.tsv'):
            rows=list(csv.DictReader(io.StringIO(path.read_text(encoding='utf-8-sig')),delimiter='\t' if suffix=='.tsv' else ','))
        elif suffix=='.xlsx':
            attachments._check_archive(path)
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
            try:
                sheet=book.active
                if (sheet.max_row or 0)>10001 or (sheet.max_column or 0)>20:raise ValueError('Task spreadsheet exceeds 10,000 rows / 20 columns')
                values=sheet.iter_rows(values_only=True);headers=[str(v) for v in next(values)];rows=[dict(zip(headers,row)) for row in values if any(v is not None for v in row)]
            finally:book.close()
        elif suffix in ('.txt','.md'):
            rows=[{'prompt':line.strip()} for line in path.read_text(encoding='utf-8-sig').splitlines() if line.strip()]
        else:raise ValueError('Import text, Markdown, CSV, TSV, JSON or XLSX tasks')
        if not isinstance(rows,list) or not 1<=len(rows)<=10000:raise ValueError('Import 1–10,000 task records')
        # Validate the entire import before mutating; batches remain atomic at the
        # plan level by requiring the import target to be paused/draft.
        state=self.owned(project,conversation)['state']
        if state not in ('draft','paused'):raise Denied('Pause the project before importing a task list')
        jobs=[]
        for index,row in enumerate(rows):
            if isinstance(row,str):row={'prompt':row}
            if not isinstance(row,dict):raise ValueError('Task rows must be text or objects')
            if row.get('recipe'):
                product={k:v for k,v in row.items() if k in ('id','recipe','title','category','price_cents','currency','pricing_evidence','purpose','sections','sheets') and v is not None}
                product.setdefault('id','import-'+record['sha256'][:8]+'-'+str(index+1));product['price_cents']=int(product.get('price_cents',0))
                if product['recipe'] not in (set(RECIPES)|DOCUMENT_RECIPES|CUSTOM_RECIPES):raise ValueError('Unsupported recipe in import')
                for stage,needs in [('generate',[]),('validate',['generate']),('package',['generate','validate'])]:jobs.append({'key':product['id']+':'+stage,'kind':stage,'payload':{'product':product},'needs':[product['id']+':'+n for n in needs]})
            else:
                key=str(row.get('key') or 'import-'+record['sha256'][:8]+'-'+str(index+1));payload=row.get('payload') or {'prompt':row.get('prompt') or row.get('task') or row.get('title','')}
                if isinstance(payload,str):payload=json.loads(payload)
                needs=row.get('needs') or []
                if isinstance(needs,str):needs=[n.strip() for n in needs.split(';') if n.strip()]
                jobs.append({'key':key,'kind':row.get('kind') or 'reason','payload':payload,'needs':needs})
        # Use one atomic validation/insertion call for a bounded import. Normal
        # model calls remain 200 tasks; this trusted parser handles large files.
        return self.add(project,conversation,jobs,importing=True)

    def scope_manifest(self,project,conversation,store,products,expires=None):
        self.owned(project,conversation)
        if store not in self.adapters:raise Denied('No supported configured store adapter; browser account actions retain existing approvals')
        if not isinstance(products,list) or not 1<=len(products)<=1000:raise ValueError('Select a bounded product batch')
        manifest={}
        for key in sorted(set(products)):
            package=self.result(project,conversation,key+':package')
            if package['state']!='completed' or not package['result'].get('qa_verified'):raise Denied('Only completed, verified packages can enter a launch scope')
            result=package['result'];listing=result['listing']
            manifest[key]={'sha256':result['sha256'],'price_cents':listing['price_cents'],'currency':listing['currency'],'title':listing['title'],'files':result['files']}
        expires=expires or time.time()+3600
        if not time.time()<expires<=time.time()+86400:raise ValueError('Launch authorization must expire within 24 hours')
        return {'project':project,'store':store,'products':manifest,'expires':expires,'allowed_changes':['create_listing','upload_exact_package','publish','verify'],'spending_limit_cents':0,'excluded':['purchases','payment_details','legal_terms','unrelated_content']}

    def approve_scope(self,manifest):
        # Host/UI calls this only after genuine approval; never a model capability.
        scope=ident()
        with self.lock:
            self.db.execute('INSERT INTO scopes VALUES(?,?,?,?,?,0)',(scope,manifest['project'],manifest['store'],encoded(manifest),manifest['expires']));self.event(manifest['project'],None,'batch_authorized',{'scope':scope,'store':manifest['store'],'products':list(manifest['products'])});self.db.commit()
        return scope

    def revoke_scope(self,project,conversation):
        self.owned(project,conversation)
        with self.lock:self.db.execute('UPDATE scopes SET revoked=1 WHERE project=?',(project,));self.event(project,None,'scope_revoked',{});self.db.commit()

    def authorization(self,project,store,package):
        for scope in self.rows('SELECT * FROM scopes WHERE project=? AND store=? AND revoked=0 AND expires>?',(project,store,time.time())):
            manifest=json.loads(scope['manifest']);approved=manifest['products'].get(package['product']);listing=package['listing']
            if approved=={'sha256':package['sha256'],'price_cents':listing['price_cents'],'currency':listing['currency'],'title':listing['title'],'files':package['files']}:return scope['id']
        raise Denied('This product, file hash, price or storefront is outside the approved launch batch')

    def checkpoint(self,job):
        if self.closed.is_set():raise InterruptedError('Project runner stopped')
        p=self.rows('SELECT state,version FROM projects WHERE id=?',(job['project'],))[0]
        current=self.rows('SELECT state FROM jobs WHERE id=?',(job['id'],))[0]['state']
        if p['state']!='running' or p['version']!=job['version'] or current not in ('running','waiting_ai'):raise InterruptedError('Project paused, cancelled or changed')

    def execute(self,job):
        project=job['project'];p=self.rows('SELECT * FROM projects WHERE id=?',(project,))[0];conversation=p['conversation'];payload=json.loads(job['payload']);product=payload.get('product',{});key=product.get('id','')
        # Aggregate cap bounds all project artifacts, independent of the normal
        # direct-file interface and protected credential directories.
        size=0
        for folder in self.root.glob('*/artifacts'):
            for path in folder.rglob('*'):
                if path.is_symlink():raise Denied('Linked artifact workspace entries are forbidden')
                if path.is_file():size+=path.stat().st_size
                if size>2*1024*1024*1024:raise Denied('Aggregate project artifact quota exceeded (2 GB)')
        factory=ProductFactory(self.root/project/'artifacts');cancel=lambda:self.checkpoint(job);cancel()
        result=lambda stage:self.result(project,conversation,key+':'+stage)['result']
        if job['kind']=='generate':return factory.generate(product,cancel)
        if job['kind']=='validate':return factory.validate(product,result('generate'),cancel)
        if job['kind']=='package':return factory.package(product,result('generate'),result('validate'),cancel)
        adapter=self.adapters.get(payload.get('store'))
        if adapter is None:raise Denied('Configure a supported store adapter before this task can run')
        package=result('package');archive=factory.path(package['archive'])
        if digest(archive)!=package['sha256']:raise Denied('Delivery file changed after QA/launch approval')
        effect_key=hashlib.sha256((project+':'+key+':'+payload['store']).encode()).hexdigest();fingerprint=hashlib.sha256(encoded(package).encode()).hexdigest()
        if job['kind']=='verify':
            actual=adapter.lookup(effect_key);adapter.verify(actual,package)
            return {'verified':True,'listing':actual,'product':key}
        cancel()
        # Read-only lookup always precedes a write, including uncertain failures.
        existing=adapter.lookup(effect_key)
        if existing:
            adapter.verify(existing,package)
            with self.lock:
                self.db.execute("UPDATE effects SET state='verified',result=?,updated=? WHERE key=?",(encoded(existing),time.time(),effect_key));self.event(project,job['id'],'external_reconciled',{'idempotency_key':effect_key});self.db.commit()
            return {'reconciled':True,'listing':existing,'product':key}
        self.authorization(project,payload['store'],package)
        effects=self.rows('SELECT * FROM effects WHERE key=?',(effect_key,))
        if effects and effects[0]['fingerprint']!=fingerprint:raise Denied('Launch content changed; do not overwrite an earlier external effect')
        with self.lock:
            self.db.execute('INSERT OR REPLACE INTO effects VALUES(?,?,?,?,?,?,?,?)',(effect_key,project,job['id'],payload['store'],'attempting',fingerprint,'{}',time.time()));self.event(project,job['id'],'external_attempt',{'idempotency_key':effect_key});self.db.commit()
        self.authorization(project,payload['store'],package);cancel()
        actual=adapter.publish(effect_key,package,archive);adapter.verify(actual,package)
        with self.lock:self.db.execute("UPDATE effects SET state='verified',result=?,updated=? WHERE key=?",(encoded(actual),time.time(),effect_key));self.db.commit()
        return {'listing':actual,'product':key,'verified':True}

    def start(self):
        with self.lock:
            if self.thread is None:
                self.thread=threading.Thread(target=self.loop,name='project-scheduler',daemon=True);self.thread.start()

    def tick(self):
        with self.lock:
            for future,job in list(self.futures.items()):
                if not future.done():continue
                self.futures.pop(future)
                try:
                    value=future.result();self.checkpoint(job)
                    if len(encoded(value))>100000:raise ValueError('Task result exceeds checkpoint bound')
                    self.db.execute("UPDATE jobs SET state='completed',result=?,error='' WHERE id=?",(encoded(value),job['id']));self.event(job['project'],job['id'],'completed',{'kind':job['kind']})
                except InterruptedError:
                    state='pending' if job['kind'] in SAFE else 'review'
                    self.db.execute("UPDATE jobs SET state=?,error='Interrupted; completed effects were not replayed' WHERE id=? AND state='running'",(state,job['id']))
                except Exception as error:
                    # Only deterministic/read-only operations get unattended retries.
                    retry=job['kind'] in SAFE|{'verify'} and job['attempts']<job['max_attempts'] and not isinstance(error,(Denied,ValueError,AssertionError))
                    state='pending' if retry else 'blocked' if isinstance(error,Denied) else 'review' if job['kind']=='publish' else 'failed'
                    reason=type(error).__name__+': '+str(error)[:500]
                    self.db.execute('UPDATE jobs SET state=?,error=?,next_at=? WHERE id=?',(state,reason,time.time()+min(30,2**job['attempts']),job['id']));self.event(job['project'],job['id'],state,{'reason':reason})
                self.db.commit()
                p=self.db.execute('SELECT conversation FROM projects WHERE id=?',(job['project'],)).fetchone()
                try:self.emit('project_progress',p['conversation'],self.summary(job['project'],p['conversation'],limit=1))
                except Exception:pass
            # Dependency failures remain explicit and do not consume workers.
            self.db.execute("UPDATE jobs SET state='pending',error='' WHERE state='blocked' AND error='A prerequisite is blocked, failed or cancelled' AND NOT EXISTS(SELECT 1 FROM dependencies d JOIN jobs j ON j.id=d.needs WHERE d.job=jobs.id AND j.state!='completed')")
            self.db.execute("UPDATE jobs SET state='blocked',error='A prerequisite is blocked, failed or cancelled' WHERE state='pending' AND EXISTS(SELECT 1 FROM dependencies d JOIN jobs j ON j.id=d.needs WHERE d.job=jobs.id AND j.state IN ('blocked','failed','cancelled','review'))")
            ready=[dict(r) for r in self.db.execute("SELECT j.* FROM jobs j JOIN projects p ON p.id=j.project WHERE j.state='pending' AND p.state='running' AND j.version=p.version AND j.next_at<=? AND NOT EXISTS(SELECT 1 FROM dependencies d JOIN jobs n ON n.id=d.needs WHERE d.job=j.id AND n.state!='completed') ORDER BY j.priority DESC,j.rowid LIMIT ?",(time.time(),max(0,self.workers-len(self.futures))))]
            for job in ready:
                if job['kind']=='reason':
                    if not self.dispatch:continue
                    if self.db.execute("SELECT 1 FROM jobs WHERE state='waiting_ai'").fetchone():continue
                    self.db.execute("UPDATE jobs SET state='waiting_ai',attempts=attempts+1 WHERE id=?",(job['id'],));self.db.commit()
                    try:self.dispatch(job)
                    except Exception as error:self.db.execute("UPDATE jobs SET state='blocked',error=? WHERE id=?",(str(error)[:500],job['id']))
                    continue
                self.db.execute("UPDATE jobs SET state='running',attempts=attempts+1 WHERE id=?",(job['id'],));job['attempts']+=1;self.db.commit()
                future=self.pool.submit(self.execute,job);self.futures[future]=job
                future.add_done_callback(lambda _:self.wake.set())
            self.db.execute("UPDATE projects SET state='completed',updated=? WHERE state='running' AND EXISTS(SELECT 1 FROM jobs WHERE project=projects.id) AND NOT EXISTS(SELECT 1 FROM jobs WHERE project=projects.id AND state!='completed')",(time.time(),))
            self.db.execute("UPDATE projects SET state='blocked' WHERE state='running' AND EXISTS(SELECT 1 FROM jobs WHERE project=projects.id AND state!='completed') AND NOT EXISTS(SELECT 1 FROM jobs WHERE project=projects.id AND state IN ('pending','running','waiting_ai'))")
            self.db.execute("INSERT OR IGNORE INTO reports SELECT id,'pending' FROM projects WHERE state='completed'");self.db.commit()
            for report in self.db.execute("SELECT p.* FROM reports r JOIN projects p ON p.id=r.project WHERE r.state='pending' LIMIT 10").fetchall():
                try:
                    self.emit('project_completed',report['conversation'],self.summary(report['id'],report['conversation'],limit=1))
                    self.db.execute("UPDATE reports SET state='delivered' WHERE project=?",(report['id'],));self.db.commit()
                except Exception:pass  # Durable outbox retries report delivery only.

    def loop(self):
        while not self.closed.is_set():
            try:self.tick()
            except Exception:
                # A scheduler fault must be visible, not a tight retry loop.
                with self.lock:
                    self.db.execute("UPDATE projects SET state='paused' WHERE state='running'");self.event(None,None,'scheduler_fault',{'reason':'Scheduler stopped dispatching; inspect local checkpoint'});self.db.commit()
            self.wake.wait(.2);self.wake.clear()

    def stop_all(self):
        with self.lock:
            self.db.execute("UPDATE projects SET state='paused' WHERE state='running'");self.db.execute('UPDATE scopes SET revoked=1');self.event(None,None,'stop_all',{});self.db.commit()
        self.wake.set()

    def active(self):return bool(self.rows("SELECT id FROM projects WHERE state='running' AND EXISTS(SELECT 1 FROM jobs WHERE project=projects.id AND state IN ('pending','running','waiting_ai'))"))

    def close(self):
        self.closed.set();self.wake.set()
        if self.thread:self.thread.join(timeout=5)
        self.pool.shutdown(wait=True,cancel_futures=True)
        with self.lock:self.db.close()
