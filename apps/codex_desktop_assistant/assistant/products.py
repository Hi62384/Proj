"""Deterministic product recipes. No model code, formulas, paths or macros run.

Generated artifacts stay inside the project's dedicated workspace. Native
LibreOffice QA operates only on files created here, with a disposable profile.
"""
import hashlib
import json
import os
import math
import re
import subprocess
import tempfile
import threading
import time
import zipfile
from pathlib import Path

OFFICE_LOCK=threading.Lock()
RECIPES={
    'cashflow':('Cash Flow Planner',['Opening cash','Cash receipts','Operating payments','Capital purchases'],[5000,12000,8000,1000],
                [('Net movement','=B6-B7-B8',3000),('Closing cash','=B5+B10',8000)],'Plan cash receipts and payments for one period; compare closing cash across scenarios.'),
    'breakeven':('Break-even Calculator',['Unit selling price','Variable cost per unit','Fixed costs','Expected unit sales'],[50,20,3000,150],
                 [('Contribution per unit','=B5-B6',30),('Break-even units','=IF(B10>0,ROUNDUP(B7/B10,0),"No break-even")',100),('Operating profit','=B8*B10-B7',1500)],'Estimate the unit sales required to cover fixed costs, and profit at a chosen sales volume.'),
    'inventory':('Inventory Reorder Planner',['Daily unit demand','Lead time in days','Safety stock units','Units available'],[12,7,30,80],
                 [('Reorder point','=B5*B6+B7',114),('Suggested reorder units','=MAX(0,B10-B8)',34)],'Calculate a reorder threshold and the gap to that threshold; this is not an economic-order-quantity model.'),
    'pricing':('Service Pricing Calculator',['Billable hours','Hourly delivery cost','Other delivery costs','Target gross margin'],[10,40,100,.4],
               [('Total delivery cost','=B5*B6+B7',500),('Required quote','=IF(B8<1,B10/(1-B8),"Margin must be below 100%")',833.3333333333334),('Gross profit','=IF(ISNUMBER(B11),B11-B10,0)',333.33333333333337)],'Build a quote from delivery costs and a target gross margin; margin is not markup.'),
    'project':('Project Budget Tracker',['Budgeted labor hours','Budgeted hourly rate','Other budgeted costs','Actual project spend'],[80,50,1200,4900],
               [('Total budget','=B5*B6+B7',5200),('Budget remaining','=B10-B8',300),('Budget used','=IF(B10>0,B8/B10,0)',.9423076923076923)],'Compare actual project spend with the labor-and-cost budget.'),
}
DOCUMENT_RECIPES={'guide','checklist'}
CUSTOM_RECIPES={'spreadsheet'}


def calculate(operation,values):
    from decimal import Decimal
    values=[Decimal(str(v or 0)) for v in values]
    if not values:raise ValueError('Calculation needs numeric inputs')
    if operation=='sum':return float(sum(values))
    if operation=='difference' and len(values)==2:return float(values[0]-values[1])
    if operation=='product':
        total=Decimal(1)
        for value in values:total*=value
        return float(total)
    if operation=='divide' and len(values)==2:return float(values[0]/values[1]) if values[1] else 0.
    if operation=='min':return float(min(values))
    if operation=='max':return float(max(values))
    raise ValueError('Unsupported declarative calculation')


def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class ProductFactory:
    def __init__(self,root,quota=512*1024*1024):
        self.root=Path(root).resolve();self.root.mkdir(parents=True,exist_ok=True)
        self.quota=quota;self.lock=threading.RLock()

    def path(self,relative):
        p=self.root/relative
        if p.is_symlink() or not p.resolve().is_relative_to(self.root):raise ValueError('Artifact path escapes the bounded workspace')
        for parent in p.parents:
            if parent==self.root:break
            if parent.is_symlink():raise ValueError('Linked artifact directories are forbidden')
        return p

    def budget(self):
        size=0
        for p in self.root.rglob('*'):
            if p.is_symlink():raise ValueError('Linked workspace entry')
            if p.is_file():size+=p.stat().st_size
        if size>self.quota:raise ValueError('Project workspace quota exceeded')
        return size

    def generate(self,product,cancel):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment,Font,PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        from docx import Document
        recipe=product['recipe']
        if recipe in DOCUMENT_RECIPES:return self.document(product,cancel)
        if recipe in CUSTOM_RECIPES:return self.spreadsheet(product,cancel)
        if recipe not in RECIPES:raise ValueError('Unsupported product recipe; use a reasoning/design task for other formats')
        cancel();self.budget()
        folder=self.path(product['id']+'/customer');folder.mkdir(parents=True,exist_ok=True)
        title,labels,defaults,metrics,purpose=RECIPES[recipe]
        title=product.get('title') or title
        workbook=Workbook();sheet=workbook.active;sheet.title='Calculator';sheet.sheet_view.showGridLines=False
        sheet.merge_cells('A1:D1');sheet['A1']=title;sheet['A1'].data_type='s';sheet['A1'].font=Font(name='Calibri',size=20,bold=True,color='FFFFFF');sheet['A1'].fill=PatternFill('solid',fgColor='134E4A');sheet.row_dimensions[1].height=36
        sheet.merge_cells('A2:D2');sheet['A2']=purpose;sheet['A2'].alignment=Alignment(wrap_text=True,vertical='center');sheet.row_dimensions[2].height=48
        sheet['A4']='EDIT BLUE INPUT CELLS';sheet['A4'].font=Font(bold=True,color='134E4A')
        for row,(label,number) in enumerate(zip(labels,defaults),5):
            sheet.cell(row,1,label);cell=sheet.cell(row,2,number);cell.font=Font(color='1565C0',size=12);cell.fill=PatternFill('solid',fgColor='E7F2FC');cell.number_format='#,##0.00';sheet.row_dimensions[row].height=26
            valid=DataValidation(type='decimal',operator='between',formula1=0,formula2=.999999 if recipe=='pricing' and row==8 else 1000000000);valid.errorTitle='Input outside supported range';valid.error='Use a nonnegative input within the stated range.';valid.showErrorMessage=True;valid.errorStyle='stop';valid.allow_blank=True;sheet.add_data_validation(valid);valid.add(cell)
        if recipe=='pricing':sheet['B8'].number_format='0.0%'
        for row,(label,formula,_) in enumerate(metrics,10):
            sheet.cell(row,1,label);sheet.cell(row,2,formula);sheet.cell(row,2).number_format='#,##0.00';sheet.cell(row,2).font=Font(bold=True,size=13,color='134E4A');sheet.row_dimensions[row].height=28
        if recipe=='project':sheet['B12'].number_format='0.0%'
        sheet['A15']='Inputs are illustrative, not market research. Change them for your business.'
        sheet.merge_cells('A15:D15');sheet['A15'].alignment=Alignment(wrap_text=True);sheet.row_dimensions[15].height=34
        sheet.column_dimensions['A'].width=35;sheet.column_dimensions['B'].width=25;sheet.column_dimensions['C'].width=8;sheet.column_dimensions['D'].width=8
        sheet.freeze_panes='B5';sheet.print_area='A1:D15';sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=1
        guide=workbook.create_sheet('Read me');guide.column_dimensions['A'].width=100
        for row,text in enumerate([title,purpose,'1. Save a working copy.','2. Edit blue cells on Calculator; amounts must use one consistent currency and period.','3. Read the calculated green outputs. Compare scenarios in copies of the workbook.','4. Blank numeric inputs are treated as zero by Excel. Negative inputs are unsupported.','5. Zero contribution cannot break even. A 100% margin cannot produce a finite quote.','No macros, external links, hidden sheets or customer data are included.','Illustrative planning tool; not tax, legal, investment or accounting advice.'],1):
            guide.cell(row,1,text);guide.cell(row,1).alignment=Alignment(wrap_text=True,vertical='top');guide.row_dimensions[row].height=42 if row>1 else 32
        guide['A1'].data_type='s';guide['A1'].font=Font(size=18,bold=True,color='134E4A');guide.print_area='A1:A9';guide.sheet_properties.pageSetUpPr.fitToPage=True;guide.page_setup.fitToWidth=1;guide.page_setup.fitToHeight=1
        target=folder/'product.xlsx';workbook.save(target)
        doc=Document();doc.add_heading(title,0);doc.add_paragraph(purpose)
        for text in ['Save an original copy before entering business data.','Edit blue input cells; formulas recalculate when opened in Excel or LibreOffice.','Use a consistent currency and reporting period. Review assumptions before relying on an output.','Consult Read me for supported inputs and edge cases.']:
            doc.add_paragraph(text,style='List Number')
        doc.add_heading('License and limitations',1);doc.add_paragraph('You may use and adapt this template within your business. Redistribution or resale of the original package is not permitted. This template is supplied as an illustrative planning aid, without a promise of business results. It is not legal, tax or investment advice. Review this proposed license before selling the product.');doc.save(folder/'Instructions.docx')
        (folder/'README.txt').write_text(title+'\n\n'+purpose+'\n\nOpen product.xlsx in Excel or LibreOffice. See Instructions.docx and the Read me worksheet.\n',encoding='utf-8')
        cancel();self.budget()
        return {'product':product['id'],'workbook':str(target.relative_to(self.root)),'sha256':digest(target),'recipe':recipe,'purpose':purpose,'title':title}

    def spreadsheet(self,product,cancel):
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        sheets=product.get('sheets',[])
        if not isinstance(sheets,list) or not 1<=len(sheets)<=3:raise ValueError('Custom workbooks require 1–3 declared worksheets')
        book=Workbook();book.remove(book.active);checks=[];formula_cells=[]
        for specification in sheets:
            cancel();name=specification.get('name','')
            if not re.fullmatch(r'[A-Za-z][A-Za-z0-9 _-]{0,24}',name) or name in book.sheetnames:raise ValueError('Invalid or duplicate sheet name')
            rows=specification.get('rows',[])
            if not rows or len(rows)>200 or any(not isinstance(r,list) or len(r)>20 for r in rows):raise ValueError('Use 1–200 rows, at most 20 columns per custom sheet')
            sheet=book.create_sheet(name);sheet.sheet_view.showGridLines=False
            for r,row in enumerate(rows,1):
                for c,value in enumerate(row,1):
                    if value is not None and not isinstance(value,(str,int,float,bool)):raise ValueError('Cells must be ordinary typed values')
                    if isinstance(value,str) and len(value)>500:raise ValueError('Cell text too long')
                    if isinstance(value,(int,float)) and (not math.isfinite(value) or abs(value)>1e12):raise ValueError('Invalid numeric input')
                    cell=sheet.cell(r,c,value)
                    if isinstance(value,str):cell.data_type='s'
                    if r==1:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='134E4A')
                    elif isinstance(value,(int,float)):cell.font=Font(color='1565C0');cell.fill=PatternFill('solid',fgColor='E7F2FC');cell.number_format='#,##0.00'
                    cell.alignment=Alignment(wrap_text=True,vertical='center')
                sheet.row_dimensions[r].height=30 if r==1 else 26
            calculations=specification.get('calculations',[])
            if not isinstance(calculations,list) or len(calculations)>100:raise ValueError('At most 100 declared calculations per sheet')
            occupied=set()
            for calculation in calculations:
                target=calculation.get('cell','');references=calculation.get('inputs',[]);operation=calculation.get('operation')
                if not re.fullmatch(r'[A-T][1-9][0-9]{0,2}',target) or not isinstance(references,list) or not 1<=len(references)<=40 or any(not isinstance(ref,str) or not re.fullmatch(r'[A-T][1-9][0-9]{0,2}',ref) for ref in references):raise ValueError('Use bounded same-sheet cell references')
                if target in references or target in occupied or sheet[target].value is not None:raise ValueError('Calculations cannot overwrite input cells or themselves')
                values=[sheet[ref].value for ref in references]
                # Repeated formulas may refer to earlier declared outputs only.
                expected_values=[next((x['expected'] for x in reversed(checks) if x['sheet']==name and x['cell']==ref),value) for ref,value in zip(references,values)]
                if any(v is not None and not isinstance(v,(int,float)) for v in expected_values):raise ValueError('Calculation input must be numeric or a preceding output')
                expected=calculate(operation,expected_values)
                if not math.isfinite(expected):raise ValueError('Calculation output is outside the supported numeric range')
                if operation=='sum':formula='=SUM('+','.join(references)+')'
                elif operation=='product':formula='=PRODUCT('+','.join(references)+')'
                elif operation in ('min','max'):formula='='+operation.upper()+'('+','.join(references)+')'
                elif operation=='difference':formula='='+references[0]+'-'+references[1]
                else:formula='=IF('+references[1]+'=0,0,'+references[0]+'/'+references[1]+')'
                sheet[target]=formula;sheet[target].font=Font(bold=True,color='134E4A');sheet[target].number_format='#,##0.00';occupied.add(target)
                checks.append({'sheet':name,'cell':target,'expected':expected});formula_cells.append({'sheet':name,'cell':target,'formula':formula})
            from openpyxl.utils import get_column_letter
            for column in range(1,max(2,sheet.max_column)+1):sheet.column_dimensions[get_column_letter(column)].width=24 if column>1 else 36
            sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
            sheet.print_area=sheet.dimensions;sheet.page_setup.orientation='landscape';sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
        folder=self.path(product['id']+'/customer');folder.mkdir(parents=True,exist_ok=True);target=folder/'product.xlsx';book.save(target);book.close()
        title=product.get('title','Business workbook');purpose=product.get('purpose','Editable business workbook with declared calculations.')
        (folder/'README.txt').write_text(title+'\n\n'+purpose+'\n\nEdit blue numeric cells; green calculated cells contain formulas. Zero denominators produce zero. Preserve headings and use consistent units. Customer data is not included. Review assumptions before business use.\n',encoding='utf-8')
        (folder/'LICENSE.txt').write_text('Proposed license: internal business use and adaptation, no redistribution of the original package. Illustrative planning aid, no promise of results; not professional advice. Owner review before real release.\n',encoding='utf-8')
        return {'product':product['id'],'workbook':str(target.relative_to(self.root)),'sha256':digest(target),'recipe':'spreadsheet','purpose':purpose,'title':title,'checks':checks,'formulas':formula_cells,'sheet_names':[s['name'] for s in sheets]}

    def document(self,product,cancel):
        from docx import Document
        from docx.shared import Pt
        cancel();self.budget();sections=product.get('sections',[])
        if not isinstance(sections,list) or not 1<=len(sections)<=25:raise ValueError('Document recipes require 1–25 substantive sections')
        folder=self.path(product['id']+'/customer');folder.mkdir(parents=True,exist_ok=True)
        doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(11)
        title=product.get('title','Business guide');doc.add_heading(title,0)
        purpose=product.get('purpose','');doc.add_paragraph(purpose)
        for section in sections:
            heading=section.get('heading','')
            if not isinstance(heading,str) or not heading.strip() or len(heading)>150:raise ValueError('Each document section needs a short heading')
            doc.add_heading(heading,1)
            paragraphs=section.get('paragraphs',[]);items=section.get('items',[])
            if not paragraphs and not items:raise ValueError('Empty filler sections are not a product')
            for text in paragraphs:
                if not isinstance(text,str) or len(text)>4000:raise ValueError('Invalid bounded document paragraph')
                doc.add_paragraph(text)
            for text in items:
                if not isinstance(text,str) or len(text)>1000:raise ValueError('Invalid bounded checklist item')
                doc.add_paragraph(('☐ ' if product['recipe']=='checklist' else '')+text,style='List Bullet')
        target=folder/'product.docx';doc.save(target)
        (folder/'README.txt').write_text(title+'\n\n'+purpose+'\n\nEditable Word document plus PDF. Review factual claims and adapt to your business.\n',encoding='utf-8')
        (folder/'LICENSE.txt').write_text('Proposed license: use and adapt within your business; no redistribution or resale of the original package. No warranty of business results. Not legal, tax, financial or other professional advice. Owner review is required before real release.\n',encoding='utf-8')
        return {'product':product['id'],'workbook':str(target.relative_to(self.root)),'sha256':digest(target),'recipe':product['recipe'],'purpose':purpose,'title':title,'headings':[s['heading'] for s in sections]}

    def office(self,source,out,format,cancel):
        executable=Path(os.environ.get('ProgramFiles','C:/Program Files'))/'LibreOffice/program/soffice.com'
        if not executable.is_file():raise RuntimeError('Install LibreOffice for native formula recalculation and workbook previews; generation alone is not validated delivery')
        while not OFFICE_LOCK.acquire(timeout=.1):cancel()
        try:
            cancel();out.mkdir(parents=True,exist_ok=True)
            target=out/(source.stem+('.xlsx' if format.startswith('xlsx') else '.pdf'))
            target.unlink(missing_ok=True)  # Never accept an earlier attempt's stale output.
            # LibreOffice's Windows profile internals exceed MAX_PATH when the
            # profile is nested below a long project/task workspace.
            with tempfile.TemporaryDirectory(prefix='codex-office-') as profile:
                command=[str(executable),'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--nologo','--nodefault','--nofirststartwizard','--convert-to',format,'--outdir',str(out),str(source)]
                process=subprocess.Popen(command,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
                deadline=time.monotonic()+90
                try:
                    while process.poll() is None:
                        cancel()
                        if time.monotonic()>deadline:raise TimeoutError('Native document validation timed out')
                        time.sleep(.1)
                    if process.returncode:raise RuntimeError('Native document conversion failed (exit '+str(process.returncode)+')')
                finally:
                    if process.poll() is None:
                        subprocess.run([str(Path(os.environ.get('WINDIR','C:/Windows'))/'System32/taskkill.exe'),'/PID',str(process.pid),'/T','/F'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0),timeout=10)
                        process.wait(timeout=10)
            if not target.is_file():raise RuntimeError('Native conversion produced no output')
            return target
        finally:OFFICE_LOCK.release()

    def validate(self,product,generated,cancel):
        if product['recipe'] in DOCUMENT_RECIPES:return self.validate_document(product,generated,cancel)
        if product['recipe'] in CUSTOM_RECIPES:return self.validate_custom(product,generated,cancel)
        from openpyxl import load_workbook
        workbook=self.path(generated['workbook'])
        if digest(workbook)!=generated['sha256']:raise ValueError('Generated workbook changed before QA')
        internal=self.path(product['id']+'/internal');internal.mkdir(parents=True,exist_ok=True)
        formula_book=load_workbook(workbook,data_only=False,keep_links=False)
        if formula_book.vba_archive or formula_book._external_links:raise ValueError('Macros/external links are not supported')
        metrics=RECIPES[product['recipe']][3]
        assert all(formula_book['Calculator'].cell(10+i,2).value==formula for i,(_,formula,_) in enumerate(metrics))
        formula_book.close()
        native=self.office(workbook,internal/'recalculated','xlsx',cancel)
        cached=load_workbook(native,data_only=True,keep_links=False)
        actual=[cached['Calculator'].cell(10+i,2).value for i in range(len(metrics))];cached.close()
        expected=[r[2] for r in metrics]
        if any(not isinstance(a,(int,float)) or abs(a-b)>1e-7*max(1,abs(b)) for a,b in zip(actual,expected)):raise ValueError('Native spreadsheet calculation differs from independent expected outputs')
        # Native-engine boundary tests in a disposable copy; never modify the customer example.
        tests=load_workbook(workbook);sheet=tests['Calculator']
        for row in range(5,9):sheet.cell(row,2,0)
        edge=internal/'edge.xlsx';tests.save(edge);tests.close()
        edge_native=self.office(edge,internal/'edge-results','xlsx',cancel)
        values=load_workbook(edge_native,data_only=True);edge_values=[values['Calculator'].cell(10+i,2).value for i in range(len(metrics))];values.close()
        edge_expected=[0,'No break-even',0] if product['recipe']=='breakeven' else [0]*len(metrics)
        if edge_values!=edge_expected:raise ValueError('Zero-input edge case failed: '+repr(edge_values))
        if product['recipe']=='pricing':
            boundary=load_workbook(workbook);boundary['Calculator']['B8']=1;boundary.save(edge);boundary.close()
            edge_native=self.office(edge,internal/'margin-results','xlsx',cancel)
            values=load_workbook(edge_native,data_only=True);assert values['Calculator']['B11'].value=='Margin must be below 100%';values.close()
        pdf=self.office(native,internal/'preview','pdf',cancel)
        import pypdfium2 as pdfium
        document=pdfium.PdfDocument(str(pdf));previews=[]
        try:
            if len(document)<2:raise ValueError('Both workbook sheets must render')
            for index in range(len(document)):
                page=document[index];bitmap=page.render(scale=1.25);image=bitmap.to_pil();target=internal/f'preview-{index+1}.png';image.save(target);image.close();bitmap.close();page.close();previews.append(str(target.relative_to(self.root)))
        finally:document.close()
        cancel();self.budget()
        return {'verified':True,'native_engine':'LibreOffice','expected':expected,'actual':actual,'zero_input_test':True,'previews':previews,'source_sha256':generated['sha256'],'native_workbook':str(native.relative_to(self.root)),'native_sha256':digest(native),'preview_hashes':{r:digest(self.path(r)) for r in previews}}

    def validate_custom(self,product,generated,cancel):
        from openpyxl import load_workbook
        source=self.path(generated['workbook'])
        if digest(source)!=generated['sha256']:raise ValueError('Custom workbook changed before QA')
        internal=self.path(product['id']+'/internal');native=self.office(source,internal/'recalculated','xlsx',cancel)
        formulas=load_workbook(native,data_only=False,keep_links=False);values=load_workbook(native,data_only=True,keep_links=False)
        try:
            if formulas._external_links or formulas.vba_archive:raise ValueError('No macros or external workbook links are supported')
            for check in generated['checks']:
                value=values[check['sheet']][check['cell']].value
                if not isinstance(value,(int,float)) or abs(value-check['expected'])>1e-7*max(1,abs(check['expected'])):raise ValueError('Native result differs from independent decimal arithmetic')
            for sheet in values:
                for row in sheet:
                    if any(cell.data_type=='e' for cell in row):raise ValueError('Native spreadsheet error')
            for check in generated['formulas']:
                if not formulas[check['sheet']][check['cell']].value.startswith('='):raise ValueError('Export lost an advertised formula')
        finally:formulas.close();values.close()
        pdf=self.office(native,internal/'pdf','pdf',cancel)
        import pypdfium2 as pdfium
        document=pdfium.PdfDocument(str(pdf));previews=[]
        try:
            if len(document)>30:raise ValueError('Workbook print layout exceeds 30 pages')
            for index in range(len(document)):
                cancel();page=document[index];bitmap=page.render(scale=1.25);image=bitmap.to_pil();target=internal/f'preview-{index+1}.png';image.save(target);image.close();bitmap.close();page.close();previews.append(str(target.relative_to(self.root)))
        finally:document.close()
        return {'verified':True,'native_engine':'LibreOffice','checks':generated['checks'],'previews':previews,'source_sha256':generated['sha256'],'native_workbook':str(native.relative_to(self.root)),'native_sha256':digest(native),'preview_hashes':{r:digest(self.path(r)) for r in previews},'edge_case_limit':'Declared supported operations use zero-safe division; bespoke business assumptions still need review'}

    def validate_document(self,product,generated,cancel):
        from docx import Document
        from pypdf import PdfReader
        source=self.path(generated['workbook'])
        if digest(source)!=generated['sha256']:raise ValueError('Document changed before QA')
        doc=Document(source);headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
        if headings!=generated['headings']:raise ValueError('Document sections differ from advertised content')
        internal=self.path(product['id']+'/internal');pdf=self.office(source,internal/'pdf','pdf',cancel)
        reader=PdfReader(pdf)
        if not 1<=len(reader.pages)<=30:raise ValueError('Document must render to 1–30 pages')
        text='\n'.join(page.extract_text() or '' for page in reader.pages)
        if any(heading not in text for heading in headings):raise ValueError('Rendered PDF is missing a section heading')
        import pypdfium2 as pdfium
        rendered=pdfium.PdfDocument(str(pdf));previews=[]
        try:
            for index in range(len(rendered)):
                cancel();page=rendered[index];bitmap=page.render(scale=1.25);image=bitmap.to_pil();target=internal/f'preview-{index+1}.png';image.save(target);image.close();bitmap.close();page.close();previews.append(str(target.relative_to(self.root)))
        finally:rendered.close()
        return {'verified':True,'native_engine':'LibreOffice','previews':previews,'source_sha256':generated['sha256'],'native_workbook':generated['workbook'],'pdf':str(pdf.relative_to(self.root)),'pdf_sha256':digest(pdf),'native_sha256':digest(source),'preview_hashes':{r:digest(self.path(r)) for r in previews},'sections_verified':headings}

    def package(self,product,generated,qa,cancel):
        if not qa.get('verified') or qa['source_sha256']!=generated['sha256']:raise ValueError('A verified unchanged product is required')
        folder=self.path(product['id']+'/customer');workbook=self.path(generated['workbook'])
        if digest(workbook) not in (generated['sha256'],digest(self.path(qa['native_workbook']))):raise ValueError('Product changed after QA')
        if not qa.get('native_sha256') or digest(self.path(qa['native_workbook']))!=qa['native_sha256']:raise ValueError('Native QA output changed or requires revalidation after upgrade')
        for relative in qa['previews']:
            if digest(self.path(relative))!=qa.get('preview_hashes',{}).get(relative):raise ValueError('Rendered preview changed after QA')
        if qa.get('pdf') and digest(self.path(qa['pdf']))!=qa.get('pdf_sha256'):raise ValueError('Rendered document changed after QA')
        import shutil
        if self.path(qa['native_workbook'])!=workbook:shutil.copyfile(self.path(qa['native_workbook']),workbook)
        if qa.get('pdf'):shutil.copyfile(self.path(qa['pdf']),folder/'product.pdf')
        from PIL import Image,ImageDraw,ImageFont
        cover=Image.new('RGB',(1400,900),'#134e4a');draw=ImageDraw.Draw(cover)
        fontpath=Path('C:/Windows/Fonts/segoeui.ttf');font=ImageFont.truetype(str(fontpath),52) if fontpath.is_file() else ImageFont.load_default(size=40)
        words=generated['title'].split();lines=[];line=''
        for word in words:
            candidate=(line+' '+word).strip()
            if draw.textlength(candidate,font=font)>1200:lines.append(line);line=word
            else:line=candidate
        lines.append(line);draw.multiline_text((85,100),'\n'.join(lines),fill='white',font=font,spacing=14)
        draw.text((85,700),'Editable files • Instructions included',fill='#99f6e4',font=ImageFont.truetype(str(fontpath),32) if fontpath.is_file() else font);cover.save(folder/'Cover.png');cover.close()
        for index,relative in enumerate(qa['previews'],1):shutil.copyfile(self.path(relative),folder/f'Preview-{index}.png')
        listing={'title':generated['title'],'description':generated['purpose'],'category':product.get('category','Small business'),'price_cents':product.get('price_cents',0),'currency':product.get('currency','USD'),'pricing_evidence':product.get('pricing_evidence','Not researched; synthetic/example price only'),'features':['Editable numeric inputs','Formula outputs','Native zero-input QA','Instructions and two worksheet previews'],'visibility':'unpublished'}
        if product['recipe'] in DOCUMENT_RECIPES:listing['features']=['Editable Word document','PDF edition','Verified section headings','Rendered page previews']
        if product['recipe'] in CUSTOM_RECIPES:listing['features']=['Editable workbook','Declared bounded calculations','Native recalculation against independent decimal arithmetic','Rendered worksheet previews']
        internal=self.path(product['id']+'/internal');(internal/'listing.json').write_text(json.dumps(listing,indent=2),encoding='utf-8')
        files={p.name:{'sha256':digest(p),'bytes':p.stat().st_size} for p in folder.iterdir() if p.is_file()}
        archive=self.path(product['id']+'/delivery.zip')
        with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as zipped:
            for name in sorted(files):zipped.write(folder/name,name)
        with zipfile.ZipFile(archive) as zipped:
            if zipped.testzip() is not None or set(zipped.namelist())!=set(files):raise ValueError('Delivery archive failed integrity check')
        cancel();self.budget()
        return {'product':product['id'],'archive':str(archive.relative_to(self.root)),'sha256':digest(archive),'bytes':archive.stat().st_size,'files':files,'listing':listing,'qa_verified':True,'visual_review':'Rendered all workbook sheets; human/model visual review still required before real release'}
