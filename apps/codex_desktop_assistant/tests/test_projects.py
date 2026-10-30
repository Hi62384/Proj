"""Durability, concurrency, scope and duplicate-effect regression tests."""
import json
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import Mock

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from assistant.projects import ProjectQueue,encoded
from assistant.permissions import Denied,Permissions
from assistant.products import ProductFactory
from assistant.attachments import Attachments
from assistant.storage import Store
from assistant.project_control import ProjectControl


class ProjectTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.q=ProjectQueue(self.root)
        self.p=self.q.create('owner','Create useful products')['project']
    def tearDown(self):self.q.close();self.tmp.cleanup()
    def product(self,index=1):return {'id':'test-'+str(index),'recipe':'cashflow','title':'Test cash flow','price_cents':1900,'currency':'USD'}
    def plan(self,count=1):
        for batch in range(0,count,40):
            products=[self.product(i) for i in range(batch,min(count,batch+40))]
            if count>1:
                for product in products:product.update(recipe='spreadsheet',sheets=[{'name':'Scenario','rows':[['Scenario','Input'],[product['id'],10]],'calculations':[]}])
            self.q.products(self.p,'owner',products)
    def wait(self,predicate,seconds=10):
        deadline=time.monotonic()+seconds
        while time.monotonic()<deadline:
            if predicate():return
            time.sleep(.02)
        self.fail('Scheduler deadline')
    def finish_package(self):
        self.plan();row=self.q.rows("SELECT * FROM jobs WHERE kind='package'")[0]
        package={'product':'test-0','sha256':'abc','listing':{'title':'Test','price_cents':1900,'currency':'USD'},'files':{'x':{'sha256':'123','bytes':1}},'qa_verified':True}
        with self.q.lock:self.q.db.execute("UPDATE jobs SET state='completed',result=? WHERE id=?",(encoded(package),row['id']));self.q.db.commit()
        return package
    def test_cycle_and_unknown_dependency_are_atomic(self):
        for jobs in ([{'key':'a','kind':'reason','needs':['b']},{'key':'b','kind':'reason','needs':['a']}],[{'key':'a','kind':'reason','needs':['missing']}]):
            with self.assertRaises(ValueError):self.q.add(self.p,'owner',jobs)
        self.assertEqual(self.q.summary(self.p,'owner')['total'],0)
    def test_deduplicated_plan_and_conflicting_reuse(self):
        self.plan();self.plan();self.assertEqual(self.q.summary(self.p,'owner')['total'],3)
        changed=self.product(0);changed['title']='Different'
        with self.assertRaises(ValueError):self.q.products(self.p,'owner',[changed])
    def test_other_chat_cannot_control_or_read_project(self):
        with self.assertRaises(Denied):self.q.summary(self.p,'other')
        with self.assertRaises(Denied):self.q.control(self.p,'other','start')
    def test_path_traversal_and_model_code_are_rejected(self):
        for value in ('../outside','a/b','C:\\private'):
            product=self.product();product['id']=value
            with self.assertRaises(ValueError):self.q.products(self.p,'owner',[product])
        with self.assertRaises(ValueError):self.q.add(self.p,'owner',[{'key':'evil','kind':'shell','payload':{'command':'whoami'}}])
        factory=ProductFactory(self.root/'bounded')
        with self.assertRaises(ValueError):factory.path('../outside')
    def test_native_generation_is_real_and_title_cannot_inject_formula(self):
        from openpyxl import load_workbook
        product=self.product();product['title']='=HYPERLINK("https://example.org","x")'
        result=ProductFactory(self.root/'artifacts').generate(product,lambda:None)
        book=load_workbook(self.root/'artifacts'/result['workbook']);self.assertEqual(book['Calculator']['A1'].data_type,'s');self.assertEqual(book['Calculator']['B10'].value,'=B6-B7-B8');self.assertTrue(book['Calculator'].data_validations.count);book.close()
    def test_one_hundred_products_bounded_workers_dependency_order(self):
        self.plan(100);running=0;peak=0;seen=set();lock=threading.Lock()
        def execute(job):
            nonlocal running,peak
            with lock:
                running+=1;peak=max(peak,running)
                dependencies=self.q.rows('SELECT needs FROM dependencies WHERE job=?',(job['id'],))
                self.assertTrue(all(d['needs'] in seen for d in dependencies))
            time.sleep(.003)
            with lock:seen.add(job['id']);running-=1
            return {'synthetic_scaling_checkpoint':True}
        self.q.execute=execute;self.q.control(self.p,'owner','start')
        self.wait(lambda:self.q.summary(self.p,'owner')['state']=='completed',30)
        self.assertEqual(len(seen),300);self.assertLessEqual(peak,2);self.assertGreaterEqual(peak,1)
        self.assertEqual(self.q.summary(self.p,'owner')['counts'],{'completed':300})
    def test_restart_requeues_safe_work_but_not_uncertain_publication(self):
        self.plan();product=self.product(0)
        self.q.add(self.p,'owner',[{'key':'external','kind':'publish','payload':{'product':product,'store':'test'}}])
        with self.q.lock:self.q.db.execute("UPDATE jobs SET state='running'");self.q.db.commit()
        self.q.close();self.q=ProjectQueue(self.root)
        states={r['kind']:r['state'] for r in self.q.rows('SELECT kind,state FROM jobs')}
        self.assertEqual(states['generate'],'pending');self.assertEqual(states['publish'],'review')
    def test_stop_and_objective_revision_revoke_launch_scope(self):
        package=self.finish_package();self.q.adapters['local']=Mock()
        scope=self.q.scope_manifest(self.p,'owner','local',['test-0']);self.q.approve_scope(scope)
        self.assertTrue(self.q.authorization(self.p,'local',package));self.q.stop_all()
        with self.assertRaises(Denied):self.q.authorization(self.p,'local',package)
        self.q.approve_scope(scope);self.q.revise(self.p,'owner','Changed objective')
        with self.assertRaises(Denied):self.q.authorization(self.p,'local',package)
        self.assertEqual(self.q.summary(self.p,'owner')['state'],'paused')
    def test_scope_is_hash_store_price_currency_and_expiration_bound(self):
        package=self.finish_package();self.q.adapters['local']=Mock();scope=self.q.scope_manifest(self.p,'owner','local',['test-0']);self.q.approve_scope(scope)
        for key,value in [('sha256','changed'),('product','other')]:
            changed=json.loads(encoded(package));changed[key]=value
            with self.assertRaises(Denied):self.q.authorization(self.p,'local',changed)
        changed=json.loads(encoded(package));changed['listing']['price_cents']=1901
        with self.assertRaises(Denied):self.q.authorization(self.p,'local',changed)
        with self.assertRaises(Denied):self.q.authorization(self.p,'another-store',package)
        with self.q.lock:self.q.db.execute('UPDATE scopes SET expires=0');self.q.db.commit()
        with self.assertRaises(Denied):self.q.authorization(self.p,'local',package)
    def test_unsupported_store_does_not_become_authorized_from_url(self):
        with self.assertRaises(Denied):self.q.scope_manifest(self.p,'owner','https://payhip.com',['x'])
    def test_import_large_csv_is_durable_and_deduplicated(self):
        attachments=Attachments(self.root/'state');source=self.root/'tasks.csv';source.write_text('task\n'+'\n'.join('Research segment '+str(i) for i in range(1000)),encoding='utf-8')
        record=attachments.add(source);self.q.import_attachment(self.p,'owner',record,attachments);self.q.import_attachment(self.p,'owner',record,attachments)
        self.assertEqual(self.q.summary(self.p,'owner')['total'],1000)
        self.q.close();self.q=ProjectQueue(self.root);self.assertEqual(self.q.summary(self.p,'owner',900,100)['next_offset'],None)
    def test_import_invalid_record_does_not_partially_apply(self):
        attachments=Attachments(self.root/'state');source=self.root/'tasks.json';source.write_text(json.dumps([{'key':'good','kind':'reason'},{'key':'bad','kind':'shell'}]))
        with self.assertRaises(ValueError):self.q.import_attachment(self.p,'owner',attachments.add(source),attachments)
        self.assertEqual(self.q.summary(self.p,'owner')['total'],0)
    def test_cancel_is_terminal_and_completed_results_retained(self):
        self.finish_package();self.q.control(self.p,'owner','cancel')
        self.assertEqual(self.q.result(self.p,'owner','test-0:package')['state'],'completed')
        with self.assertRaises(Denied):self.q.control(self.p,'owner','resume')
    def test_no_automatic_retry_of_uncertain_publication(self):
        product=self.product();self.q.add(self.p,'owner',[{'key':'publish','kind':'publish','payload':{'product':product,'store':'local'}}])
        self.q.execute=Mock(side_effect=TimeoutError('Response lost after external attempt'));self.q.control(self.p,'owner','start')
        self.wait(lambda:self.q.summary(self.p,'owner')['counts'].get('review')==1)
        self.assertEqual(self.q.execute.call_count,1)
    def test_launch_denial_never_stores_scope(self):
        self.finish_package();store=Store(self.root/'owner-state');owner=store.new();permissions=Permissions();permissions.resume(owner)
        controller=ProjectControl(permissions,store,Attachments(store.directory),Mock(return_value=False),Mock());controller._queue=self.q
        with self.q.lock:self.q.db.execute('UPDATE projects SET conversation=? WHERE id=?',(owner,self.p));self.q.db.commit()
        self.q.adapters['local']=Mock()
        with self.assertRaises(Denied):controller.call(owner,'authorize_launch',{'project':self.p,'store':'local','products':['test-0']})
        self.assertFalse(self.q.rows('SELECT * FROM scopes'));store.close()

    def test_declarative_workbook_emits_only_host_defined_calculations(self):
        from openpyxl import load_workbook
        product={'id':'custom','recipe':'spreadsheet','title':'Sales planner','sheets':[{'name':'Sales','rows':[['Units','Price','Revenue'],[10,25,None]],'calculations':[{'cell':'C2','operation':'product','inputs':['A2','B2']}]}]}
        factory=ProductFactory(self.root/'custom');result=factory.generate(product,lambda:None)
        self.assertEqual(result['checks'][0]['expected'],250)
        workbook=load_workbook(factory.path(result['workbook']));self.assertEqual(workbook['Sales']['C2'].value,'=PRODUCT(A2,B2)');workbook.close()
        product['sheets'][0]['calculations'][0]['operation']='WEBSERVICE'
        with self.assertRaises(ValueError):factory.generate(product,lambda:None)
    def test_declarative_references_cannot_escape_sheet_or_replace_input(self):
        factory=ProductFactory(self.root/'custom')
        for target,reference in [('C2','[private.xlsx]Sheet1!A1'),('A2','A2')]:
            product={'id':'custom','recipe':'spreadsheet','sheets':[{'name':'Sales','rows':[['Units'],[10]],'calculations':[{'cell':target,'operation':'sum','inputs':[reference]}]}]}
            with self.assertRaises(ValueError):factory.generate(product,lambda:None)
    def test_document_recipe_contains_real_sections_and_checklist(self):
        from docx import Document
        product={'id':'launch-checklist','recipe':'checklist','title':'Launch checklist','sections':[{'heading':'Prepare files','items':['Validate delivery files','Check preview accuracy']} ]}
        factory=ProductFactory(self.root/'doc');result=factory.generate(product,lambda:None)
        document=Document(factory.path(result['workbook']));text='\n'.join(p.text for p in document.paragraphs)
        self.assertIn('Prepare files',text);self.assertIn('Validate delivery files',text)
    def test_synthetic_scope_cannot_be_reused_for_a_different_recipe(self):
        package=self.finish_package();store=Store(self.root/'scope-state');owner=store.new();permissions=Permissions();permissions.resume(owner)
        controller=ProjectControl(permissions,store,Attachments(store.directory),Mock(),Mock());controller._queue=self.q
        product=self.product(0);manifest={'project':self.p,'store':'synthetic','products':{'test-0':{'title':product['title'],'price_cents':1900,'currency':'USD'}}}
        synthetic={'store':'synthetic','products':[product]}
        self.assertTrue(controller.synthetic_allowed(synthetic,manifest))
        changed=dict(product,recipe='guide')
        with self.q.lock:self.q.db.execute("UPDATE jobs SET payload=? WHERE project=? AND key='test-0:generate'",(encoded({'product':changed}),self.p));self.q.db.commit()
        self.assertFalse(controller.synthetic_allowed(synthetic,manifest));store.close()
    def test_package_rejects_changed_native_output_or_preview_after_qa(self):
        import shutil
        from assistant.products import digest
        factory=ProductFactory(self.root/'integrity');product=self.product();generated=factory.generate(product,lambda:None)
        native=factory.path('test-1/internal/native.xlsx');native.parent.mkdir(parents=True);shutil.copyfile(factory.path(generated['workbook']),native)
        preview=factory.path('test-1/internal/preview.png');preview.write_bytes(b'original synthetic preview')
        qa={'verified':True,'source_sha256':generated['sha256'],'native_workbook':str(native.relative_to(factory.root)),'native_sha256':digest(native),'previews':[str(preview.relative_to(factory.root))],'preview_hashes':{str(preview.relative_to(factory.root)):digest(preview)}}
        native.write_bytes(b'changed after native QA')
        with self.assertRaises(ValueError):factory.package(product,generated,qa,lambda:None)
        shutil.copyfile(factory.path(generated['workbook']),native);preview.write_bytes(b'changed after rendering')
        with self.assertRaises(ValueError):factory.package(product,generated,qa,lambda:None)


if __name__=='__main__':unittest.main()
