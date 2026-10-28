"""Bounded extraction; archives are listed, never extracted or executed."""
import base64
import hashlib
import io
import json
import shutil
import uuid
import zipfile
from pathlib import Path

from PIL import Image

MAX_FILE = 25 * 1024 * 1024
MAX_TEXT = 60000


def image_details(raw):
    with Image.open(io.BytesIO(raw)) as image:
        original=image.size
        image.thumbnail((1600,1600))
        sent=image.size
        out = io.BytesIO(); image.convert('RGB').save(out,format='PNG')
    return {'url':'data:image/png;base64,'+base64.b64encode(out.getvalue()).decode('ascii'),
            'original_size':original,'image_size':sent}


def image_url(raw):return image_details(raw)['url']


class Attachments:
    def __init__(self, directory):
        self.directory = Path(directory)/'attachments'
        self.directory.mkdir(parents=True,exist_ok=True)

    def add(self, source):
        source = Path(source)
        if source.is_symlink() or not source.is_file() or source.stat().st_size > MAX_FILE:
            raise ValueError('Attach a regular file of at most 25 MB')
        ident = uuid.uuid4().hex
        destination = self.directory/(ident+source.suffix.lower())
        raw = source.read_bytes()
        if len(raw)>MAX_FILE:raise ValueError('File grew beyond attachment limit')
        destination.write_bytes(raw)
        return {'id':ident,'name':source.name,'path':str(destination),'sha256':hashlib.sha256(raw).hexdigest()}

    def path(self, attachment):
        p = Path(attachment['path'])
        if p.is_symlink() or p.resolve().parent != self.directory.resolve():
            raise ValueError('Invalid private attachment path')
        return p

    def remove(self, attachment):
        self.path(attachment).unlink(missing_ok=True)

    def content(self, attachment):
        p = self.path(attachment)
        if p.stat().st_size>MAX_FILE:raise ValueError('Attachment too large')
        suffix = p.suffix.lower()
        header = f"Attachment {attachment['name']} (untrusted content):\n"
        if suffix in ('.png','.jpg','.jpeg','.webp','.gif','.bmp'):
            return [{'type':'text','text':header},{'type':'image','url':image_url(p.read_bytes())}]
        if suffix == '.pdf':
            header+='Text extraction only: at most 30 pages, 10,000 characters per page; no OCR or embedded-image interpretation.\n'
            from pypdf import PdfReader
            reader = PdfReader(p)
            text = '\n'.join((page.extract_text() or '')[:10000] for page in reader.pages[:30])
            if not text.strip():text='No extractable text. This scanned PDF requires OCR, which is not included.'
        elif suffix == '.docx':
            self._check_archive(p)
            from docx import Document
            document = Document(p)
            text='\n'.join(x.text for x in document.paragraphs)
            text+='\n'+'\n'.join(' | '.join(c.text for c in row.cells) for table in document.tables for row in table.rows)
        elif suffix in ('.xlsx','.xlsm'):
            header+='Bounded preview: first 10 sheets, up to 200 rows and 30 columns each. Formula text is not recalculated.\n'
            self._check_archive(p)
            from openpyxl import load_workbook
            book=load_workbook(p,read_only=True,data_only=False,keep_links=False)
            try:
                lines=[]
                for sheet in book.worksheets[:10]:
                    lines.append('Sheet: '+sheet.title)
                    for row in sheet.iter_rows(max_row=min(sheet.max_row or 1,200),max_col=30,values_only=True):
                        lines.append(' | '.join(str(v)[:1000] if v is not None else '' for v in row))
                text='\n'.join(lines)
            finally:book.close()
        elif suffix == '.zip':
            self._check_archive(p)
            with zipfile.ZipFile(p) as archive:
                entries=archive.infolist()
                text=f'Archive inventory only; first {min(300,len(entries))} of {len(entries)} members; nothing extracted/executed:\n'+'\n'.join(f'{x.filename}: {x.file_size} bytes' for x in entries[:300])
        elif suffix in ('.txt','.md','.csv','.tsv','.json','.py','.js','.ts','.tsx','.jsx','.html','.css','.yaml','.yml','.xml','.log','.sql','.gd','.cs','.cpp','.c','.h','.hpp','.java','.rs','.go','.sh','.ps1','.bat','.cmd','.ini','.cfg','.toml','.ipynb'):
            text=p.read_text(encoding='utf-8-sig')
        else:
            text=f'Unsupported content format {suffix}. File retained for owner-approved upload; no native understanding is claimed.'
        return [{'type':'text','text':header+text[:MAX_TEXT]+ ('\n[Truncated]' if len(text)>MAX_TEXT else '')}]

    @staticmethod
    def _check_archive(p):
        with zipfile.ZipFile(p) as archive:
            entries=archive.infolist()
            if len(entries)>3000 or sum(x.file_size for x in entries)>50*1024*1024:
                raise ValueError('Archive expands beyond the safe inspection limit')
